import sqlite3
import json
import os
import sys
import re
import subprocess
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Ensure UTF-8 printing in Windows terminal
if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')

def get_desktop_path():
    paths_to_try = [
        os.path.join(os.path.expanduser("~"), "Desktop"),
        os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop"),
        r"C:\Users\Abou Mohamed\Desktop",
        r"C:\Users\Abou Mohamed\OneDrive\Desktop"
    ]
    for p in paths_to_try:
        if os.path.exists(p):
            return p
    return "."

def pull_database_from_device():
    adb_path = r"C:\Users\Abou Mohamed\AppData\Local\Android\Sdk\platform-tools\adb.exe"
    if not os.path.exists(adb_path):
        adb_path = "adb"
        
    cmds = [
        [adb_path, "shell", "touch /data/local/tmp/fb_db_temp && chmod 666 /data/local/tmp/fb_db_temp"],
        [adb_path, "shell", "run-as com.example.light_app_controles cp databases/abdelwahab-jemla-com-default-rtdb.europe-west1.firebasedatabase.app_default cache/fb_db_temp"],
        [adb_path, "shell", "run-as com.example.light_app_controles cp cache/fb_db_temp /data/local/tmp/fb_db_temp"],
        [adb_path, "pull", "/data/local/tmp/fb_db_temp", "fb_db_temp"]
    ]
    
    print("Pulling latest database from connected device...")
    for cmd in cmds:
        try:
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=10)
            if res.returncode != 0:
                print(f"ADB warning: {' '.join(cmd)} failed. Stderr: {res.stderr.strip()}")
                return False
        except Exception as e:
            print(f"ADB exception: {e}")
            return False
    print("Database pulled successfully from device.")
    return True

def export_records_to_sheet(ws, records, dynamic_keys, sheet_title):
    # Set sheet view to show grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Stylings
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    
    border_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Order columns alphabetically
    columns = sorted(list(dynamic_keys))
    
    # Write Header Row
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    # Write Data Rows
    for row_idx, rec in enumerate(records, 2):
        for col_idx, col_name in enumerate(columns, 1):
            val = rec.get(col_name, "")
            if isinstance(val, (list, dict)):
                val = json.dumps(val, ensure_ascii=False)
            elif isinstance(val, bool):
                val = "Vrai" if val else "Faux"
                
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = left_align

    # Set row heights
    ws.row_dimensions[1].height = 28
    for r in range(2, len(records) + 2):
        ws.row_dimensions[r].height = 20
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 50)

def main():
    # Attempt to pull latest database
    pull_database_from_device()

    # Find the database file to read from
    db_options = ["fb_db_temp", "fb_db", "fb_db_93f88"]
    db_path = None
    for option in db_options:
        if os.path.exists(option):
            db_path = option
            break
            
    if not db_path:
        print("Error: None of the Firebase offline cache databases (fb_db_temp, fb_db, fb_db_93f88) were found in the root directory.")
        return

    print(f"Reading Firebase cache from SQLite file: '{db_path}'")
    
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # 1. Query M13Tariffication
    m13_records = []
    m13_keys = set()
    try:
        cursor.execute("SELECT path, value FROM serverCache WHERE path LIKE '%/M13Tariffication/%'")
        rows_m13 = cursor.fetchall()
        for path, val_blob in rows_m13:
            try:
                val_str = val_blob.decode('utf-8', errors='ignore')
                val_json = json.loads(val_str)
                if isinstance(val_json, dict):
                    m13_records.append(val_json)
                    for k in val_json.keys():
                        m13_keys.add(k)
            except Exception:
                pass
    except Exception as e:
        print(f"Error querying M13Tariffication: {e}")
        
    # 2. Query M03Couleur
    m3_records = []
    m3_keys = set()
    try:
        cursor.execute("SELECT path, value FROM serverCache WHERE path LIKE '%/M03Couleur/%'")
        rows_m3 = cursor.fetchall()
        for path, val_blob in rows_m3:
            # Skip list keys summaries
            if "-00_listKeys" in path:
                continue
            try:
                val_str = val_blob.decode('utf-8', errors='ignore')
                val_json = json.loads(val_str)
                if isinstance(val_json, dict):
                    # Filter to only keep records where count_Don_Depot > 0
                    if val_json.get("count_Don_Depot", 0) > 0:
                        m3_records.append(val_json)
                        for k in val_json.keys():
                            m3_keys.add(k)
            except Exception:
                pass
    except Exception as e:
        print(f"Error querying M03Couleur: {e}")
        
    conn.close()
    
    print(f"Found {len(m13_records)} M13Tariffication records and {len(m3_records)} M03Couleur records.")
    
    if not m13_records and not m3_records:
        print("No M13 or M3 records found. Export aborted.")
        return
        
    # Write to Excel
    wb = openpyxl.Workbook()
    
    # Sheet 1: M13Tariffication
    if m13_records:
        ws_m13 = wb.active
        ws_m13.title = "M13Tariffication"
        export_records_to_sheet(ws_m13, m13_records, m13_keys, "M13Tariffication")
        # If we have M3, create second sheet
        if m3_records:
            ws_m3 = wb.create_sheet(title="M03Couleur")
            export_records_to_sheet(ws_m3, m3_records, m3_keys, "M03Couleur")
    elif m3_records:
        ws_m3 = wb.active
        ws_m3.title = "M03Couleur"
        export_records_to_sheet(ws_m3, m3_records, m3_keys, "M03Couleur")
        
    # Save Workbook to Desktop
    desktop = get_desktop_path()
    excel_filename = "fb_m13_m3_ref.xlsx"
    dest_path = os.path.join(desktop, excel_filename)
    
    try:
        wb.save(dest_path)
        print(f"Excel file successfully saved to: {dest_path}")
    except Exception as e:
        dest_path = excel_filename
        wb.save(dest_path)
        print(f"Failed to write to desktop. Excel file saved to fallback path: {dest_path}. Error: {e}")
        
    # Write Markdown Report
    report_lines = [
        "# Firebase M13 & M3 References Export Report",
        f"\n- **Database Source**: `{db_path}`",
        f"- **Desktop Path**: [{dest_path}](file:///{dest_path.replace(os.sep, '/')})",
        f"- **Timestamp**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- **M13 Tariffication Records**: {len(m13_records)}",
        f"- **M03 Couleur Records**: {len(m3_records)}",
        "\n## Preview of M13 Tariffication Records (First 5):\n",
        "| KeyID | DebugInfos | PrixCurrency | TypeChoisi | ProductKeyID |",
        "| --- | --- | --- | --- | --- |"
    ]
    
    for rec in m13_records[:5]:
        r_key = rec.get("keyID", "")
        r_deb = rec.get("debugInfos", "")
        r_prix = rec.get("prixCurrency", "")
        r_type = rec.get("typeChoisi", "")
        r_prod = rec.get("parent_M1Produit_KeyId", "")
        report_lines.append(f"| `{r_key}` | {r_deb} | {r_prix} | {r_type} | `{r_prod}` |")
        
    report_lines.extend([
        "\n## Preview of M03 Couleur Records (First 5):\n",
        "| KeyID | NomCouleurStrSiSonImageDispo | ParentBProduitInfosKeyID | CreationTimestamp |",
        "| --- | --- | --- | --- |"
    ])
    
    for rec in m3_records[:5]:
        r_key = rec.get("keyID", "")
        r_nom = rec.get("nomCouleurStrSiSonImageDispo", "")
        r_par = rec.get("parentBProduitInfosKeyID", "")
        r_time = rec.get("creationTimestamp", "")
        report_lines.append(f"| `{r_key}` | {r_nom} | `{r_par}` | {r_time} |")
        
    report_path = "copy_/fb_db/last_fb_m13_m3_excel.md"
    os.makedirs(os.path.dirname(report_path), exist_ok=True)
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(report_lines))
        
    # Write last_query_m13_m3.md log
    query_log = f"""# Last Firebase M13 & M3 Excel Export Info

- **Action**: Exported Firebase M13 & M3 to Excel
- **Timestamp**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
- **Exported File**: {dest_path}
- **M13 Records**: {len(m13_records)}
- **M3 Records**: {len(m3_records)}
"""
    log_dir = "app/src/main/java/skill_agent/fb_db"
    os.makedirs(log_dir, exist_ok=True)
    with open(os.path.join(log_dir, "last_query_m13_m3.md"), "w", encoding="utf-8") as f:
        f.write(query_log)
        
    print("M13 & M3 references Excel export task complete.")

if __name__ == "__main__":
    main()
