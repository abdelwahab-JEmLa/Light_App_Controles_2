import sqlite3
import json
import os
import sys
import re
import subprocess
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# Ensure UTF-8 printing in Windows terminal
if sys.platform.startswith('win'):
    sys.stdout.reconfigure(encoding='utf-8')

# Standard columns from M01Produit.csv
STANDARD_COLUMNS = [
    "id", "keyID", "creationTimestamp", "bsonObjectId", 
    "dernierTimeTampsSynchronisationAvecFireBase", "dernierFireBaseUpdateTimestamps", 
    "count_Don_Depot", "processPositioningInFactory", "idParentCategorie", 
    "positionDonSonCesFrereCategorieProduits", "nom", "nomMutable", "nomArab", 
    "autreNomDarticle", "etateActuelleOnFusionAvecBaseDonne", "tag_prioriter_str", 
    "nombreUniteInt", "nombreProduitDonSonCarton", "its_Carton", "cartonState", 
    "heldPrioriteDemandAuGrossist", "position_store_3jamale", 
    "dernier_timeTamps_position_store_3jamale", "prixDefiniParGerant", "prixVent", 
    "cachePrixVent", "pourcentage_Prix_Progressive", "prixAchat", 
    "prixAchatDernierTimeTempUpdate", "clientPrixVentUnite", "afficheUniteAuPrint", 
    "actualiseSonImage", "actualiseSonImageTest2", "afficheCesDetailPourComptBsonId", 
    "disponibilityEtates", "disponibilityEtates_Pour_presentaion_par_Camion", 
    "keyFireBase", "couleur1", "couleur2", "couleur3", "couleur4", "couleur5", 
    "couleur6", "couleur7", "couleur8", "couleur9", "idcolor1", "idcolor2", 
    "idcolor3", "idcolor4", "idcolor5", "idcolor6", "idcolor7", "idcolor8", 
    "idcolor9", "nomCategorie2", "affichageUniteState", "commmentSeVent", 
    "afficheBoitSiUniter", "minQuan", "monBenfice", "neaon2", 
    "funChangeImagsDimention", "nomCategorie", "neaon1", "lastUpdateState", 
    "dateCreationCategorie", "prixDeVentTotaleChezClient", "benficeTotaleEntreMoiEtClien", 
    "benificeTotaleEn2", "monPrixAchatUniter", "monPrixVentUniter", 
    "articleHaveUniteImages", "itsNewArrivale", "imageDimention", "idForSearchArticles", 
    "prioriter", "quantite_Boit_Par_Carton", "setIN_Vent_Its_Quantity_Represent"
]

def get_desktop_path():
    paths_to_try = [
        os.path.join(os.path.expanduser("~"), "Desktop"),
        os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop"),
        r"C:\Users\Abou Mohamed\Desktop",
        r"C:\Users\Abou Mohamed\OneDrive\Desktop"
    ]
    for p in paths_to_try:
        if os.path.exists(p):
            return p
    return "."

def pull_database_from_device():
    adb_path = r"C:\Users\Abou Mohamed\AppData\Local\Android\Sdk\platform-tools\adb.exe"
    if not os.path.exists(adb_path):
        adb_path = "adb"
        
    cmds = [
        [adb_path, "shell", "touch /data/local/tmp/fb_db_temp && chmod 666 /data/local/tmp/fb_db_temp"],
        [adb_path, "shell", "run-as com.example.light_app_controles cp databases/abdelwahab-jemla-com-default-rtdb.europe-west1.firebasedatabase.app_default cache/fb_db_temp"],
        [adb_path, "shell", "run-as com.example.light_app_controles cp cache/fb_db_temp /data/local/tmp/fb_db_temp"],
        [adb_path, "pull", "/data/local/tmp/fb_db_temp", "fb_db_temp"]
    ]
    
    print("Pulling latest database from connected device...")
    for cmd in cmds:
        try:
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=10)
            if res.returncode != 0:
                print(f"ADB warning: {' '.join(cmd)} failed. Stderr: {res.stderr.strip()}")
                return False
        except Exception as e:
            print(f"ADB exception: {e}")
            return False
    print("Database pulled successfully from device.")
    return True

def search_codebase_for_todo():
    todo_pattern = re.compile(r'//\s*TODO:\s*(fb_m1_excel|m1_excel)', re.IGNORECASE)
    search_dir = "app/src/main/java"
    if not os.path.exists(search_dir):
        return None, None
        
    for root, dirs, files in os.walk(search_dir):
        for file in files:
            if file.endswith('.kt') or file.endswith('.java'):
                file_path = os.path.join(root, file)
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    lines = f.readlines()
                for idx, line in enumerate(lines):
                    match = todo_pattern.search(line)
                    if match:
                        return file_path, idx + 1
    return None, None

def remove_todo_comment(file_path, line_number):
    if not file_path or not os.path.exists(file_path):
        return
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        lines = f.readlines()
    
    if 1 <= line_number <= len(lines):
        comment_line = lines[line_number - 1]
        if comment_line.strip().startswith("//"):
            lines[line_number - 1] = ""
        else:
            idx = comment_line.lower().find("//")
            if idx != -1:
                lines[line_number - 1] = comment_line[:idx] + "\n"
                
        with open(file_path, 'w', encoding='utf-8') as f:
            f.writelines(lines)
        print(f"Removed TODO comment from {file_path}:{line_number}")

def main():
    source_file = "Direct Chat Prompt"
    line_number = None

    # Search codebase for TODO: fb_m1_excel or TODO: m1_excel comment
    found_file, found_line = search_codebase_for_todo()
    if found_file:
        source_file = found_file
        line_number = found_line
        print(f"Found active TODO comment at {source_file}:{line_number}")

    # Attempt to pull latest database
    pull_database_from_device()

    # Find the database file to read from
    db_options = ["fb_db_temp", "fb_db", "fb_db_93f88"]
    db_path = None
    for option in db_options:
        if os.path.exists(option):
            db_path = option
            break
            
    if not db_path:
        print("Error: None of the Firebase offline cache databases (fb_db_temp, fb_db, fb_db_93f88) were found in the root directory.")
        return

    print(f"Reading Firebase cache from SQLite file: '{db_path}'")
    
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT path, value FROM serverCache WHERE path LIKE '%M01Produit%'")
        rows = cursor.fetchall()
    except Exception as e:
        print(f"Error querying SQLite database: {e}")
        conn.close()
        return
    
    conn.close()
    
    print(f"Found {len(rows)} M01Produit records in {db_path}.")
    if not rows:
        print("No M01Produit records found. Make sure the database has been populated.")
        return
        
    records = []
    dynamic_keys = set()
    for path, val_blob in rows:
        try:
            val_str = val_blob.decode('utf-8', errors='ignore')
            val_json = json.loads(val_str)
            if isinstance(val_json, dict):
                records.append(val_json)
                for k in val_json.keys():
                    dynamic_keys.add(k)
        except Exception as e:
            pass
            
    print(f"Successfully parsed {len(records)} records.")
    
    # Determine columns list: standard columns first, followed by any other keys found in records
    columns = list(STANDARD_COLUMNS)
    for k in sorted(dynamic_keys):
        if k not in columns:
            columns.append(k)
            
    # Write to Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "M01Produit"
    
    # Set sheet view to show grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Stylings
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    
    border_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Write Header Row
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        
    # Write Data Rows
    for row_idx, rec in enumerate(records, 2):
        for col_idx, col_name in enumerate(columns, 1):
            val = rec.get(col_name, "")
            # Convert JSON lists/dicts to strings for Excel compatibility
            if isinstance(val, (list, dict)):
                val = json.dumps(val, ensure_ascii=False)
            elif isinstance(val, bool):
                val = "Vrai" if val else "Faux"
                
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            # Align text vs numeric fields
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = left_align

    # Set row heights
    ws.row_dimensions[1].height = 28
    for r in range(2, len(records) + 2):
        ws.row_dimensions[r].height = 20
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 50)
        
    # Save Workbook to Desktop
    desktop = get_desktop_path()
    excel_filename = "fb_m1_ref.xlsx"
    dest_path = os.path.join(desktop, excel_filename)
    
    try:
        wb.save(dest_path)
        print(f"Excel file successfully saved to: {dest_path}")
    except Exception as e:
        # Fallback to current directory if desktop writing fails
        dest_path = excel_filename
        wb.save(dest_path)
        print(f"Failed to write to desktop. Excel file saved to fallback path: {dest_path}. Error: {e}")
        
    # Write Markdown Report
    report_lines = [
        "# Firebase M1 References Export Report",
        f"\n- **Database Source**: `{db_path}`",
        f"- **Desktop Path**: [{dest_path}](file:///{dest_path.replace(os.sep, '/')})",
        f"- **Timestamp**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- **Total Exported Records**: {len(records)}",
        f"- **Total Columns**: {len(columns)}",
        f"- **Trigger Source**: {source_file}" + (f":{line_number}" if line_number else ""),
        "\n## Preview of First 5 Exported Records:\n",
        "| ID | KeyID | Nom | Prix Achat | Prix Vent | Categorie |",
        "| --- | --- | --- | --- | --- | --- |"
    ]
    
    for rec in records[:5]:
        r_id = rec.get("id", "")
        r_key = rec.get("keyID", "")
        r_nom = rec.get("nom", "")
        r_achat = rec.get("prixAchat", "")
        r_vent = rec.get("prixVent", "")
        r_cat = rec.get("nomCategorie", "")
        report_lines.append(f"| {r_id} | `{r_key}` | {r_nom} | {r_achat} | {r_vent} | {r_cat} |")
        
    report_path = "copy_/fb_db/last_fb_m1_excel.md"
    os.makedirs(os.path.dirname(report_path), exist_ok=True)
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("\n".join(report_lines))
        
    # Write last_query.md log
    query_log = f"""# Last Firebase M1 Excel Export Info

- **Source File**: {source_file}
- **Action**: Exported Firebase M01Produit (M1) to Excel
- **Timestamp**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
- **Exported File**: {dest_path}
- **Total Records**: {len(records)}
"""
    log_dir = "app/src/main/java/skill_agent/fb_db"
    os.makedirs(log_dir, exist_ok=True)
    with open(os.path.join(log_dir, "last_query.md"), "w", encoding="utf-8") as f:
        f.write(query_log)
        
    # Clean up codebase TODO
    if line_number:
        remove_todo_comment(source_file, line_number)
        
    print("M1 references Excel export task complete.")

if __name__ == "__main__":
    main()
