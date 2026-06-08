import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def get_desktop_path():
    paths_to_try = [
        os.path.join(os.path.expanduser("~"), "Desktop"),
        os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop"),
        r"C:\Users\Abou Mohamed\Desktop",
        r"C:\Users\Abou Mohamed\OneDrive\Desktop"
    ]
    for p in paths_to_try:
        if os.path.exists(p):
            return p
    return "."

def main():
    desktop = get_desktop_path()
    
    # Determine which files to process
    files_to_process = []
    if len(sys.argv) >= 2 and sys.argv[1].lower() != "all":
        filename = sys.argv[1]
        # Handle .xlx vs .xlsx typos
        if filename.endswith(".xlx"):
            filename = filename + "s"
        elif not filename.endswith(".xlsx"):
            filename = filename + ".xlsx"
        files_to_process.append(filename)
    else:
        # Scan Desktop for all xlsx files, skipping temporary ones
        for f in os.listdir(desktop):
            if f.endswith(".xlsx") and not f.startswith("~$"):
                files_to_process.append(f)
                
    if not files_to_process:
        print("No Excel files found to process on the Desktop.")
        return
        
    print(f"Found {len(files_to_process)} Excel file(s) to scan: {', '.join(files_to_process)}")
    
    for filename in files_to_process:
        excel_path = os.path.join(desktop, filename)
        print(f"\nOpening Excel file: {excel_path}")
        try:
            wb = openpyxl.load_workbook(excel_path)
        except Exception as e:
            print(f"Error: Could not open {filename}. It might be locked or corrupted. Error: {e}")
            continue
            
        modified = False
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # --- 1. Custom M13Tariffication Extraction ---
            if sheet_name == "M13Tariffication":
                cell_a1 = ws.cell(1, 1)
                if cell_a1.comment:
                    text = cell_a1.comment.text.lower()
                    if any(kw in text for kw in ["extract", "separated", "fichie"]):
                        print(f"[{filename}] Found extract comment on M13Tariffication cell A1: '{cell_a1.comment.text.strip()}'")
                        dest_m13_path = os.path.join(desktop, "M13Tariffication.xlsx")
                        print(f"[{filename}] Extracting M13Tariffication sheet to {dest_m13_path}...")
                        
                        # Copy sheet to a new workbook
                        new_wb = openpyxl.Workbook()
                        new_ws = new_wb.active
                        new_ws.title = ws.title
                        new_ws.views.sheetView[0].showGridLines = True
                        
                        # Copy values and styling
                        for r in range(1, ws.max_row + 1):
                            new_ws.row_dimensions[r].height = ws.row_dimensions[r].height
                            for c in range(1, ws.max_column + 1):
                                src_cell = ws.cell(row=r, column=c)
                                new_cell = new_ws.cell(row=r, column=c, value=src_cell.value)
                                if src_cell.has_style:
                                    new_cell.font = Font(name=src_cell.font.name, size=src_cell.font.size, bold=src_cell.font.bold, italic=src_cell.font.italic, color=src_cell.font.color) if src_cell.font else None
                                    new_cell.fill = PatternFill(fill_type=src_cell.fill.fill_type, start_color=src_cell.fill.start_color, end_color=src_cell.fill.end_color) if src_cell.fill else None
                                    new_cell.alignment = Alignment(horizontal=src_cell.alignment.horizontal, vertical=src_cell.alignment.vertical, wrap_text=src_cell.alignment.wrap_text) if src_cell.alignment else None
                                    new_cell.border = Border(left=src_cell.border.left, right=src_cell.border.right, top=src_cell.border.top, bottom=src_cell.border.bottom) if src_cell.border else None
                        
                        # Copy column widths
                        for col in ws.columns:
                            col_letter = get_column_letter(col[0].column)
                            new_ws.column_dimensions[col_letter].width = ws.column_dimensions[col_letter].width
                            
                        # Remove comment from new file
                        if new_ws.cell(1, 1).comment:
                            new_ws.cell(1, 1).comment = None
                            
                        try:
                            new_wb.save(dest_m13_path)
                            print(f"[{filename}] Successfully saved M13Tariffication.xlsx to Desktop.")
                        except Exception as ex:
                            print(f"Error saving M13Tariffication.xlsx: {ex}")
                            
                        # Clear comment from original sheet
                        cell_a1.comment = None
                        modified = True
                        continue

            # --- 2. Custom M03Couleur Formula Lookup ---
            if sheet_name == "M03Couleur":
                # Check for comment in cell S1 (col 19)
                cell_s1 = ws.cell(1, 19)
                s_comment = cell_s1.comment.text.lower() if cell_s1.comment else ""
                
                if "aller au m13" in s_comment:
                    print(f"[{filename}] Found comment on M03Couleur cell S1: '{cell_s1.comment.text.strip()}'")
                    # Clear comment
                    cell_s1.comment = None
                    
                    # Update all formulas in Column S to point to external workbook M13Tariffication.xlsx
                    for r in range(2, ws.max_row + 1):
                        cell = ws.cell(row=r, column=19)
                        cell.value = f"=SUMIFS('[M13Tariffication.xlsx]M13Tariffication'!Q:Q, '[M13Tariffication.xlsx]M13Tariffication'!L:L, R{r}, '[M13Tariffication.xlsx]M13Tariffication'!T:T, \"Tariff_ItsWorkInGrossist_SuperGros\")"
                        cell.font = Font(name="Segoe UI", size=10)
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.border = Border(left=Side(style="thin", color="D3D3D3"),
                                             right=Side(style="thin", color="D3D3D3"),
                                             top=Side(style="thin", color="D3D3D3"),
                                             bottom=Side(style="thin", color="D3D3D3"))
                    
                    print(f"[{filename}] Updated all formulas in Column S of M03Couleur to point to external M13Tariffication.xlsx.")
                    modified = True
                    continue

                cell_q1 = ws.cell(1, 17)
                cell_r1 = ws.cell(1, 18)
                q_comment = cell_q1.comment.text.lower() if cell_q1.comment else ""
                r_comment = cell_r1.comment.text.lower() if cell_r1.comment else ""
                
                if ("formule" in q_comment or "tariff_itsworkingrossist_supergros" in q_comment) or "par ca" in r_comment:
                    print(f"[{filename}] Found formula comments on M03Couleur: Q1='{q_comment.strip()}', R1='{r_comment.strip()}'")
                    
                    # Clear comments from cells
                    cell_q1.comment = None
                    cell_r1.comment = None
                    
                    # Check if column already exists
                    col_exists = False
                    for col in range(1, ws.max_column + 1):
                        if ws.cell(1, col).value == "Tariff_ItsWorkInGrossist_SuperGros":
                            col_exists = True
                            break
                            
                    if not col_exists:
                        # Insert column at index 19 (Column S)
                        ws.insert_cols(19)
                        
                        # Setup Header
                        header_cell = ws.cell(row=1, column=19, value="Tariff_ItsWorkInGrossist_SuperGros")
                        header_cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                        header_cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        header_cell.border = Border(left=Side(style="thin", color="D3D3D3"),
                                                     right=Side(style="thin", color="D3D3D3"),
                                                     top=Side(style="thin", color="D3D3D3"),
                                                     bottom=Side(style="thin", color="D3D3D3"))
                        
                        # Set formulas for rows (using key from column 18/R)
                        for r in range(2, ws.max_row + 1):
                            cell = ws.cell(row=r, column=19)
                            cell.value = f'=SUMIFS(M13Tariffication!Q:Q, M13Tariffication!L:L, R{r}, M13Tariffication!T:T, "Tariff_ItsWorkInGrossist_SuperGros")'
                            cell.font = Font(name="Segoe UI", size=10)
                            cell.alignment = Alignment(horizontal="right", vertical="center")
                            cell.border = Border(left=Side(style="thin", color="D3D3D3"),
                                                 right=Side(style="thin", color="D3D3D3"),
                                                 top=Side(style="thin", color="D3D3D3"),
                                                 bottom=Side(style="thin", color="D3D3D3"))
                        
                        ws.column_dimensions[get_column_letter(19)].width = 35
                        print(f"[{filename}] Created column 'Tariff_ItsWorkInGrossist_SuperGros' at column S (19) with SUMIFS formulas.")
                    else:
                        print(f"[{filename}] Column 'Tariff_ItsWorkInGrossist_SuperGros' already exists.")
                    
                    modified = True
                    continue

            # --- 2b. Custom balance.xlsx sum of condepo m3 * Tariff_ItsWorkInGrossist_SuperGros ---
            if filename == "balance.xlsx" and sheet_name == "Feuil1":
                cell_a1 = ws.cell(1, 1)
                a1_comment = cell_a1.comment.text.lower() if cell_a1.comment else ""
                if "sum" in a1_comment and "condepo" in a1_comment and "tariff_itsworkingrossist_supergros" in a1_comment:
                    print(f"[{filename}] Found comment on cell A1: '{cell_a1.comment.text.strip()}'")
                    # Clear comment
                    cell_a1.comment = None
                    
                    # Write formula directly in A1
                    cell_a1.value = "=SUMPRODUCT('[fb_m13_m3_ref.xlsx]M03Couleur'!$E$2:$E$114, '[fb_m13_m3_ref.xlsx]M03Couleur'!$S$2:$S$114)"
                    cell_a1.font = Font(name="Segoe UI", size=11, bold=True)
                    cell_a1.alignment = Alignment(horizontal="right", vertical="center")
                    cell_a1.border = Border(left=Side(style="thin", color="D3D3D3"),
                                            right=Side(style="thin", color="D3D3D3"),
                                            top=Side(style="thin", color="D3D3D3"),
                                            bottom=Side(style="thin", color="D3D3D3"))
                    ws.column_dimensions['A'].width = 35
                    print(f"[{filename}] Wrote SUMPRODUCT formula in cell A1 and cleared comment.")
                    modified = True
                    continue

            # --- 3. Standard Constraint Filter (e.g. count_Don_Depot > 0) ---
            comment_found = False
            target_cell_coord = None
            comment_text = ""
            
            for coord, cell in list(ws._cells.items()):
                if cell.comment:
                    text = cell.comment.text
                    if "dep" in text.lower() and (">0" in text or "> 0" in text or "supérieur" in text.lower()):
                        comment_found = True
                        target_cell_coord = cell.coordinate
                        comment_text = text
                        break
                        
            if comment_found:
                print(f"[{filename}] Found comment on cell {target_cell_coord}: '{comment_text.strip()}'")
                headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
                col_idx = None
                col_name = ""
                
                for idx, h in enumerate(headers, 1):
                    if h == "count_Don_Depot":
                        col_idx = idx
                        col_name = h
                        break
                if not col_idx:
                    for idx, h in enumerate(headers, 1):
                        if h and "depot" in h.lower():
                            col_idx = idx
                            col_name = h
                            break
                if not col_idx:
                    for idx, h in enumerate(headers, 1):
                        if h and "dep" in h.lower():
                            col_idx = idx
                            col_name = h
                            break
                            
                if col_idx:
                    print(f"[{filename}] Applying constraint: {col_name} > 0")
                    rows_to_keep = []
                    deleted_count = 0
                    keep_count = 0
                    
                    for row_idx, row_values in enumerate(ws.iter_rows(values_only=True), 1):
                        if row_idx == 1:
                            continue
                        val = row_values[col_idx - 1]
                        try:
                            val_num = float(val) if val is not None else 0
                        except (ValueError, TypeError):
                            val_num = 0
                        if val_num > 0:
                            rows_to_keep.append(row_values)
                            keep_count += 1
                        else:
                            deleted_count += 1
                            
                    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                    data_font = Font(name="Segoe UI", size=10)
                    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    left_align = Alignment(horizontal="left", vertical="center")
                    border_side = Side(border_style="thin", color="D3D3D3")
                    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
                    
                    idx_pos = wb.sheetnames.index(sheet_name)
                    wb.remove(ws)
                    new_ws = wb.create_sheet(title=sheet_name, index=idx_pos)
                    new_ws.views.sheetView[0].showGridLines = True
                    
                    for col_c, col_name_val in enumerate(headers, 1):
                        cell = new_ws.cell(row=1, column=col_c, value=col_name_val)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = center_align
                        cell.border = thin_border
                        
                    for row_r, rec in enumerate(rows_to_keep, 2):
                        for col_c, val in enumerate(rec, 1):
                            cell = new_ws.cell(row=row_r, column=col_c, value=val)
                            cell.font = data_font
                            cell.border = thin_border
                            if isinstance(val, (int, float)):
                                cell.alignment = Alignment(horizontal="right", vertical="center")
                            else:
                                cell.alignment = left_align
                                
                    new_ws.row_dimensions[1].height = 28
                    for r in range(2, len(rows_to_keep) + 2):
                        new_ws.row_dimensions[r].height = 20
                        
                    for col in new_ws.columns:
                        max_len = 0
                        col_letter = get_column_letter(col[0].column)
                        for cell in col:
                            val_str = str(cell.value or '')
                            if len(val_str) > max_len:
                                max_len = len(val_str)
                        new_ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 50)
                        
                    print(f"[{filename}] Filtered sheet '{sheet_name}'. Kept {keep_count} rows, deleted {deleted_count} rows.")
                    modified = True
                else:
                    print(f"[{filename}] Warning: Could not find any column matching 'dep' in sheet '{sheet_name}' headers.")
            else:
                print(f"[{filename}] No matching constraint comments found in sheet '{sheet_name}'.")
                
        if modified:
            try:
                wb.save(excel_path)
                print(f"[{filename}] Successfully saved modifications to {excel_path}")
            except Exception as e:
                print(f"Error: Could not save modifications to {filename}. The file might be open in Excel. Error: {e}")
        else:
            print(f"[{filename}] No changes made.")

if __name__ == "__main__":
    main()
